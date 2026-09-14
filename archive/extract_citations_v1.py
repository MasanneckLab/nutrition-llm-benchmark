#!/usr/bin/env python3
"""
Citation extraction from the 180 model responses (version 1, archived).

This is the version with which the withdrawn citation counts were produced: for
Sonar Reasoning Pro it captured only the URL list and missed the bracket citations
in the running text (97 citations in 31 responses). The parsers parse_gpt and
parse_oe are imported unchanged by analysis/extract_citations.py; parse_sonar was
replaced there. Only comments and docstrings were translated.

Three system-specific parsers, because the citation formats differ structurally:

  GPT-5.5   [author, year, *title/journal*, PMID/DOI: ...]       -> bibliographic
  Open Evidence   numbered reference list:
                  N.
                  Title.
                  Journal. Year. Authors.<type tag>              -> bibliographic + own label
  Sonar Reasoning Pro   section "Sources:" with bare URLs        -> web links

Output: citations_raw.csv (one row per citation occurrence)
"""
import openpyxl, re, csv, sys, os

BASE = os.environ.get('NUTRITION_RAW', '.')  # raw workbooks, not deposited
OUT = os.path.dirname(os.path.abspath(__file__))

MODEL_NAME = {'Open Evidence': 'Open Evidence',
              'gpt-5.5-2026-04-23': 'GPT-5.5',
              'sonar-reasoning-pro': 'Sonar Reasoning Pro'}


def load_data():
    wb = openpyxl.load_workbook(f'{BASE}/27f926ae-29c1-4dfe-b5b5-bf2e24df6b96.xlsx', data_only=True)
    ws = wb['Row Key']
    h = [c.value for c in ws[1]]
    key = {dict(zip(h, r))['Rating ID']: dict(zip(h, r))
           for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}

    wb2 = openpyxl.load_workbook(f'{BASE}/f8339a6a-6c4d-43cc-bb5a-a710308d87b6.xlsx', data_only=True)
    ws2 = wb2['Rating']
    h2 = [c.value for c in ws2[1]]
    rows = [dict(zip(h2, r)) for r in ws2.iter_rows(min_row=2, values_only=True) if r[0]]
    for x in rows:
        x['Model'] = MODEL_NAME[key[x['Rating ID']]['Actual model']]
    return rows


# ---------------------------------------------------------------- GPT-5.5
GPT_CIT = re.compile(r'\[([^\[\]]{25,400}?)\]')

def parse_gpt(text):
    out = []
    for m in GPT_CIT.finditer(text):
        raw = re.sub(r'\s+', ' ', m.group(1)).strip()
        # only entries that look like a citation: contain a year
        if not re.search(r'\b(19|20)\d{2}\b', raw):
            continue
        year = re.search(r'\b(19|20)\d{2}\b', raw).group(0)
        pmid = re.search(r'PMID:?\s*(\d{6,9})', raw)
        doi = re.search(r'(10\.\d{4,9}/[^\s,\]]+)', raw)
        author = raw.split(',')[0].strip().strip('*')
        title = ''
        tm = re.search(r'\*(.+?)\*', raw)
        if tm:
            title = tm.group(1)
        else:
            parts = [p.strip() for p in raw.split(',')]
            if len(parts) >= 3:
                title = parts[2]
        out.append(dict(raw=raw, author=author, year=year, title=title,
                        pmid=pmid.group(1) if pmid else '',
                        doi=doi.group(1) if doi else '', url='', oe_tag=''))
    return out


# ---------------------------------------------- Open Evidence: reference list
OE_TAGS = ['Guideline', 'Review', 'Recent', 'Landmark', 'Highly Cited', 'Meta-Analysis', 'RCT']

def parse_oe(text):
    # remove Citavi plug-in text (interface artefact, not model output)
    text = re.sub(r'Titel anhand dieser (DOI|Pubmed-ID|PMID)[^\n]*?Citavi-Projekt übernehmen\.?', ' ', text)
    lines = [l.strip() for l in text.split('\n')]
    out = []
    i = 0
    while i < len(lines):
        if re.fullmatch(r'\d{1,3}\.', lines[i]):
            block = []
            j = i + 1
            while j < len(lines) and len(block) < 4:
                if lines[j]:
                    if re.fullmatch(r'\d{1,3}\.', lines[j]):
                        break
                    block.append(lines[j])
                j += 1
            if len(block) >= 2:
                title = block[0].rstrip('.')
                meta = block[1]
                ym = re.search(r'\b(19|20)\d{2}\b', meta)
                year = ym.group(0) if ym else ''
                journal = meta.split('.')[0].strip()
                authors = '.'.join(meta.split('.')[2:]).strip() if meta.count('.') >= 2 else ''
                tag = ''
                for t in OE_TAGS:
                    if meta.rstrip().endswith(t) or (len(block) > 2 and block[2].strip() == t):
                        tag = t
                        break
                if not tag:
                    tm = re.search(r'(' + '|'.join(OE_TAGS) + r')\s*$', meta)
                    if tm:
                        tag = tm.group(1)
                author = re.split(r'[,;]', authors)[0].strip() if authors else ''
                out.append(dict(raw=re.sub(r'\s+', ' ', ' | '.join(block))[:300],
                                author=author, year=year, title=title,
                                journal=journal, pmid='', doi='', url='', oe_tag=tag))
            i = j
        else:
            i += 1
    return out


# ------------------------------------------------------- Sonar: URL list
URL_RE = re.compile(r'https?://[^\s\)\]<>"]+')

def parse_sonar(text):
    out = []
    seen = set()
    for m in URL_RE.finditer(text):
        u = m.group(0).rstrip('.,;')
        if u in seen:
            continue
        seen.add(u)
        dom = re.sub(r'^www\.', '', u.split('/')[2].lower())
        out.append(dict(raw=u, author='', year='', title='', journal='',
                        pmid='', doi='', url=u, domain=dom, oe_tag=''))
    return out


def main():
    rows = load_data()
    recs = []
    for x in rows:
        t = str(x['Model answer'] or '')
        m = x['Model']
        if m == 'GPT-5.5':
            cits = parse_gpt(t)
        elif m == 'Open Evidence':
            cits = parse_oe(t)
        else:
            cits = parse_sonar(t)
        for c in cits:
            recs.append(dict(
                rating_id=x['Rating ID'], model=m, domain=x['Domain'],
                question_id=x['Question ID'], qtype=x['Question type'],
                item_guideline=x['Guideline'],
                accuracy=x['Accuracy rating'],
                glr=x['Guideline reference correct?'],
                author=c.get('author', ''), year=c.get('year', ''),
                title=c.get('title', ''), journal=c.get('journal', ''),
                pmid=c.get('pmid', ''), doi=c.get('doi', ''),
                url=c.get('url', ''), domain_host=c.get('domain', ''),
                oe_tag=c.get('oe_tag', ''), raw=c.get('raw', '')))

    fn = f'{OUT}/citations_raw.csv'
    with open(fn, 'w', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, fieldnames=list(recs[0].keys()))
        w.writeheader()
        w.writerows(recs)

    import collections
    per = collections.Counter(r['model'] for r in recs)
    ans = collections.defaultdict(set)
    for r in recs:
        ans[r['model']].add(r['rating_id'])
    print(f'{len(recs)} citation occurrences extracted -> {fn}')
    for m in ['Open Evidence', 'GPT-5.5', 'Sonar Reasoning Pro']:
        print(f'  {m:22s} {per[m]:5d} citations in {len(ans[m]):2d}/60 responses '
              f'(mean {per[m]/max(len(ans[m]),1):.1f} per response)')


if __name__ == '__main__':
    main()
