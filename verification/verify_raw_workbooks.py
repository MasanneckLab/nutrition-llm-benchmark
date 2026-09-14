#!/usr/bin/env python3
"""Check the deposited ratings against the private source workbooks.

The workbooks are intentionally not part of this repository. Pass their paths on
the command line when an authorized local copy is available. Rater one is Rainer
and rater two is Gero in ``data/ratings_raters.csv``.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent.parent
MODEL_NAMES = {
    "Open Evidence": "Open Evidence",
    "gpt-5.5-2026-04-23": "GPT-5.5",
    "sonar-reasoning-pro": "Sonar Reasoning Pro",
}
RATING_COLUMNS = {
    "Accuracy rating": "accuracy",
    "Guideline reference correct?": "guideline_reference",
    "Safety concern?": "safety_concern",
}


def worksheet_rows(path: Path, sheet_name: str) -> list[dict]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    worksheet = workbook[sheet_name]
    iterator = worksheet.iter_rows(values_only=True)
    headers = list(next(iterator))
    rows = [dict(zip(headers, values)) for values in iterator if values and values[0]]
    workbook.close()
    return rows


def read_raters() -> dict[str, dict[str, str]]:
    with (ROOT / "data" / "ratings_raters.csv").open(encoding="utf-8", newline="") as handle:
        return {row["rating_id"]: row for row in csv.DictReader(handle)}


def compare_book(
    name: str,
    rows: list[dict],
    deposited: dict[str, dict[str, str]],
    deposited_prefix: str,
) -> None:
    by_id = {row["Rating ID"]: row for row in rows}
    if set(by_id) != set(deposited):
        raise AssertionError(f"{name}: Rating IDs differ from the deposited table")
    for workbook_column, deposited_column in RATING_COLUMNS.items():
        mismatches = [
            rating_id
            for rating_id, row in by_id.items()
            if (row.get(workbook_column) or "")
            != deposited[rating_id][f"{deposited_prefix}_{deposited_column}"]
        ]
        if mismatches:
            raise AssertionError(
                f"{name}: {workbook_column} differs in {len(mismatches)} rows"
            )
    print(f"  {name:<12} 180 IDs and three rating fields: OK")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--key", type=Path, required=True)
    parser.add_argument("--consensus", type=Path, required=True)
    parser.add_argument("--rater-one", type=Path, required=True, help="Rainer workbook")
    parser.add_argument("--rater-two", type=Path, required=True, help="Gero workbook")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    key = {
        row["Rating ID"]: row
        for row in worksheet_rows(args.key, "Row Key")
    }
    consensus_rows = worksheet_rows(args.consensus, "Rating")
    rater_one_rows = worksheet_rows(args.rater_one, "Rating")
    rater_two_rows = worksheet_rows(args.rater_two, "Rating")
    deposited = read_raters()

    selected_ids = {row["Rating ID"] for row in consensus_rows}
    if len(selected_ids) != 180 or not selected_ids <= set(key):
        raise AssertionError("Consensus must contain 180 unique IDs present in the key")
    selected_models = {MODEL_NAMES[key[rating_id]["Actual model"]] for rating_id in selected_ids}
    if selected_models != set(MODEL_NAMES.values()):
        raise AssertionError(f"Unexpected selected systems: {selected_models}")

    print("PRIVATE WORKBOOK RECONCILIATION")
    compare_book("consensus", consensus_rows, deposited, "consensus")
    compare_book("rater one", rater_one_rows, deposited, "rater_1")
    compare_book("rater two", rater_two_rows, deposited, "rater_2")

    hash_mismatches = []
    for row in consensus_rows:
        expected = key[row["Rating ID"]].get("Model answer SHA-256")
        actual = hashlib.sha256((row.get("Model answer") or "").encode("utf-8")).hexdigest()
        if expected and expected != actual:
            hash_mismatches.append(row["Rating ID"])
    print(f"  answer hashes differing from this key version: {len(hash_mismatches)}")
    print("  evidence-quality ratings are not compared: the deposited columns are from the later rerating")
    print("\nRAW WORKBOOK CHECKS PASSED")


if __name__ == "__main__":
    main()
