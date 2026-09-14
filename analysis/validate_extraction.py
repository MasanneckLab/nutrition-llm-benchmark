#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Validation of the citation extraction against the stored response texts.

Counts per response, independently of the parser, how many citations of each
format are in the text and compares that with the number of extracted citations.
Deviations are printed. Exactly this check would have found the error in
version 1: Sonar Reasoning Pro cites in two formats, only one was captured.

The raw data are not deposited. Path via environment variable:
    export NUTRITION_RAW=/path/to/raw/data

Usage:  python3 validate_extraction.py <citations csv>
        (default ../data/citations_long_v3.csv; the archive version reports 32 deviations)
"""
import openpyxl, re, csv, sys, os, collections

RAW = os.environ.get('NUTRITION_RAW', '.')
KEY = 'Nutrition_master_rating_key_3models_DO_NOT_SHARE.xlsx'
CONSENSUS = 'Final_Konsens_Nutrition_master_rating_Konsensabgleich.xlsx'
MODEL_NAME = {'Open Evidence': 'Open Evidence', 'gpt-5.5-2026-04-23': 'GPT-5.5',
              'sonar-reasoning-pro': 'Sonar Reasoning Pro'}

URL_RE = re.compile(r'https?://[^\s\)\]<>"]+')
BR_RE = re.compile(r'\[([^\[\]]{20,400}?)\]')
YEAR_RE = re.compile(r'\b(19|20)\d{2}\b')
# Numbered list item. Date lines such as "12. Juni 2025" are excluded, otherwise
# the check counts time stamps as references.
MONTHS = (r'januar|februar|m[äa]rz|april|mai|juni|juli|august|september|oktober|november|dezember|'
          r'january|february|march|may|june|july|october|december')
NUM_ITEM = re.compile(r'(?mi)^\s*(\d{1,3})\.\s+(?!(?:%s)\b)(\S)' % MONTHS)
LABEL_PREFIX = re.compile(r'^\s*(guideline|systematic review|meta-?analysis|trial|rct|'
                          r'observational|insufficient evidence)\s*[:\-–]\s*', re.I)


def is_bibliographic(inner):
    s = LABEL_PREFIX.sub('', inner).strip()
    return bool(YEAR_RE.search(s)) and s.count(',') >= 2 and 20 <= len(s) <= 400


def answers():
    wb = openpyxl.load_workbook(os.path.join(RAW, KEY), data_only=True)
    ws = wb['Row Key']; h = [c.value for c in ws[1]]
    key = {d['Rating ID']: d for d in (dict(zip(h, r)) for r in ws.iter_rows(min_row=2, values_only=True)) if d.get('Rating ID')}
    wb2 = openpyxl.load_workbook(os.path.join(RAW, CONSENSUS), data_only=True)
    ws2 = wb2['Rating']; h2 = [c.value for c in ws2[1]]
    for r in ws2.iter_rows(min_row=2, values_only=True):
        d = dict(zip(h2, r))
        if d.get('Rating ID'):
            yield d['Rating ID'], MODEL_NAME[key[d['Rating ID']]['Actual model']], d['Model answer'] or ''


def expected(model, text):
    """Expected number of distinct citations per format, counted independently of the parser."""
    urls = {u.rstrip('.,;') for u in URL_RE.findall(text)}
    brs = set()
    for m in BR_RE.finditer(text):
        if is_bibliographic(m.group(1)):
            brs.add(re.sub(r'\s+', ' ', LABEL_PREFIX.sub('', m.group(1))).strip().lower())
    if model == 'Open Evidence':
        return {'numbered_list': len({m[0] for m in NUM_ITEM.findall(text)})}
    if model == 'GPT-5.5':
        return {'bracket': len(brs)}
    return {'url_list': len(urls), 'bracket': len(brs)}


def main(path):
    got = collections.defaultdict(collections.Counter)
    for r in csv.DictReader(open(path)):
        fmt = r.get('source_format') or ''
        if not fmt:
            # file without a format column: derive the format from the record itself
            # so that older files can be checked too
            if r.get('url'):
                fmt = 'url_list'
            elif r['model'] == 'Open Evidence':
                fmt = 'numbered_list'
            else:
                fmt = 'bracket'
        if fmt == 'both':
            got[r['rating_id']]['url_list'] += 1
            got[r['rating_id']]['bracket'] += 1
        else:
            got[r['rating_id']][fmt] += 1

    problems = collections.Counter()
    lines = []
    for rid, model, text in answers():
        exp = expected(model, text)
        have = got.get(rid, collections.Counter())
        for fmt, n_exp in exp.items():
            n_have = have.get(fmt, 0)
            if n_have != n_exp:
                d = n_have - n_exp
                problems[(model, fmt)] += 1
                lines.append((rid, model, fmt, n_exp, n_have, d))

    print('checked: %s' % path)
    print('180 responses, deviations: %d\n' % len(lines))
    if problems:
        print('%-22s %-14s %s' % ('System', 'format', 'responses with deviation'))
        for (m, f), n in sorted(problems.items()):
            print('%-22s %-14s %d' % (m, f, n))
        print()
        print('%-16s %-22s %-12s %6s %6s %6s' % ('response', 'system', 'format', 'text', 'CSV', 'diff'))
        for rid, m, f, e, h, d in lines[:25]:
            print('%-16s %-22s %-12s %6d %6d %+6d' % (rid, m, f, e, h, d))
        if len(lines) > 25:
            print('... and %d more' % (len(lines) - 25))
        fehlend = sum(-d for *_, d in lines if d < 0)
        print('\nmissing citations in total: %d' % fehlend)
    else:
        print('No deviation. Every citation recognisable in the text is in the file.')
    return len(lines)


if __name__ == '__main__':
    default = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'data', 'citations_long_v3.csv')
    print('Reference: against citations_long_v3.csv one deviation remains (a GPT-5.5 response citing the\n'
          'same guideline twice; counting convention, not a gap). Against archive/citations_long_v1.csv: 32.\n')
    sys.exit(1 if main(sys.argv[1] if len(sys.argv) > 1 else default) else 0)
