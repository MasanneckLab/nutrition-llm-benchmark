# -*- coding: utf-8 -*-
"""In-text citations: which of the extracted citations does the response actually invoke in its running text?

Background
----------
The three systems carry their sources differently:

  GPT-5.5              bibliographic bracket citations inside the sentence. Every
                       citation is in the running text by construction.
  Open Evidence        numbered reference list at the end, invoked in the text
                       as [n]. Not every entry is invoked.
  Sonar Reasoning Pro  bracket citations in the text plus a list of URLs. The
                       list does not come from the model: the Perplexity API
                       returns the retrieved sources as a separate field beside
                       the message, and the collection script appended it to the
                       text. It is invoked as [n].

For the primary citation analysis (Table 5) only what the response itself invokes
as evidence is counted. This script decides that per citation and writes an
assignment file so that table5_citation_sources.py runs without the raw data.

Rule per citation:
  source_format bracket or both    -> invoked (it is in the running text)
  numbered_list or url_list        -> invoked if its position in the list occurs
                                      in the text as [n]
  position cannot be determined    -> invoked (conservative, counted)

The raw data are not deposited:
    export NUTRITION_RAW=/path/to/raw/data

Usage:
    python3 in_text_citations.py [../data/citations_long_v3.csv] [../data/citations_in_text.csv]

Requires openpyxl.
"""
import csv, sys, os, re, collections
import openpyxl

RAW = os.environ.get('NUTRITION_RAW', '.')
KONSENS = 'Final_Konsens_Nutrition_master_rating_Konsensabgleich.xlsx'
MARK = re.compile(r'\[(\d{1,3})\](?!\()')
URL = re.compile(r'https?://[^\s\)\]<>"]+')


def antworttexte():
    ws = openpyxl.load_workbook(os.path.join(RAW, KONSENS), data_only=True)['Rating']
    kopf = [c.value for c in ws[1]]
    out = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(kopf, r))
        if d.get('Rating ID'):
            out[d['Rating ID']] = d['Model answer'] or ''
    return out


def norm(s):
    return re.sub(r'[^a-z0-9]', '', (s or '').lower())[:60]


def main(quelle, ziel):
    txt = antworttexte()
    rows = list(csv.DictReader(open(quelle, encoding='utf-8')))
    nach_antwort = collections.defaultdict(list)
    for r in rows:
        nach_antwort[r['rating_id']].append(r)

    belegt, stat = {}, collections.Counter()
    for rid, rs in nach_antwort.items():
        t = txt[rid]
        m = re.search(r'Sources:', t)
        body, liste = (t[:m.start()], t[m.end():]) if m else (t, '')
        marks = {int(g) for g in MARK.findall(body)}

        # Position in the URL list (Sonar)
        urls, gesehen = [], set()
        for u in URL.finditer(liste):
            uu = u.group(0).rstrip('.,;')
            if uu not in gesehen:
                gesehen.add(uu)
                urls.append(uu)
        url_pos = {u: i for i, u in enumerate(urls, 1)}

        # Position in the numbered reference list (Open Evidence)
        eintraege = re.findall(r'(?m)^\s*(\d{1,3})\.\s*\n(.{10,400}?)(?=\n\s*\d{1,3}\.\s*\n|\Z)', t, re.S)
        num_pos = {}
        for nummer, block in eintraege:
            num_pos.setdefault(norm(block.split('\n')[0]), int(nummer))

        for r in rs:
            fmt, mod = r['source_format'], r['model']
            if fmt in ('bracket', 'both'):
                ja, grund = True, 'bracket citation in running text'
            elif fmt == 'url_list':
                p = url_pos.get((r['url'] or '').strip().rstrip('.,;'))
                ja = p in marks
                grund = 'URL list, position %s' % (p if p else 'undetermined')
                if p is None:
                    ja, grund = True, 'URL list, position undetermined'
            else:
                k = norm(r['title'] or r['citation_raw_text'])
                p = num_pos.get(k)
                if p is None:
                    for kk, vv in num_pos.items():
                        if kk and (kk.startswith(k[:40]) or k.startswith(kk[:40])):
                            p = vv
                            break
                if p is None:
                    ja, grund = True, 'reference list, position undetermined'
                else:
                    ja, grund = p in marks, 'reference list, position %d' % p
            belegt[r['citation_id']] = (1 if ja else 0, grund)
            stat[(mod, ja)] += 1

    with open(ziel, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow(['citation_id', 'invoked_in_text', 'reason'])
        for r in rows:
            ja, grund = belegt[r['citation_id']]
            w.writerow([r['citation_id'], ja, grund])

    print('%d citations assigned, written: %s' % (len(rows), ziel))
    print()
    print('%-22s %10s %12s' % ('', 'in text', 'listed only'))
    for m in ['Open Evidence', 'GPT-5.5', 'Sonar Reasoning Pro']:
        print('%-22s %10d %12d' % (m, stat[(m, True)], stat[(m, False)]))
    print()
    print('Reference values - a run must reproduce these:')
    print('  Open Evidence          290 in text, 141 listed only')
    print('  GPT-5.5                168 in text,   0 listed only')
    print('  Sonar Reasoning Pro    400 in text, 463 listed only')


if __name__ == '__main__':
    q = sys.argv[1] if len(sys.argv) > 1 else os.path.join('..', 'data', 'citations_long_v3.csv')
    z = sys.argv[2] if len(sys.argv) > 2 else os.path.join('..', 'data', 'citations_in_text.csv')
    main(q, z)
