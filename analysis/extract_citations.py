#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Extraction of the citations from the stored model responses (version 2).

Change from version 1: Sonar Reasoning Pro cites in TWO formats. Version 1
captured only the URL list at the end of the response. Bibliographic bracket
citations in the running text, such as

    [Volkert, 2019, ESPEN guideline on clinical nutrition and hydration in
     geriatrics/Clin Nutr, 10.1016/j.clnu.2018.12.025]

were missed entirely. Version 2 captures both formats and merges duplicates: the
same document can occur in both formats and counts once.

Formats recognised per system
  Open Evidence         numbered reference list at the end            (unchanged)
  GPT-5.5               bibliographic bracket citations in the text   (unchanged)
  Sonar Reasoning Pro   URL list  AND  bracket citations in the text  (new)

The parsers for Open Evidence and GPT-5.5 are imported unchanged from
archive/extract_citations_v1.py.

The raw data are not deposited. Path via environment variable:
    export NUTRITION_RAW=/path/to/raw/data

Output: ../data/intermediate/citations_raw_v2.csv
Requires openpyxl.
"""
import openpyxl, re, csv, os, collections

RAW = os.environ.get('NUTRITION_RAW', '.')
KEY = 'Nutrition_master_rating_key_3models_DO_NOT_SHARE.xlsx'
CONSENSUS = 'Final_Konsens_Nutrition_master_rating_Konsensabgleich.xlsx'
HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, '..', 'data', 'intermediate')
ARCHIVE = os.path.join(HERE, '..', 'archive')
MODEL_NAME = {'Open Evidence': 'Open Evidence',
              'gpt-5.5-2026-04-23': 'GPT-5.5',
              'sonar-reasoning-pro': 'Sonar Reasoning Pro'}

DOI_RE = re.compile(r'10\.\d{4,9}/[^\s,;\]\)]+')
PMID_LABEL = re.compile(r'PMID[:\s]*(\d{7,8})', re.I)
PMID_BARE = re.compile(r'(?:^|[,\s])(\d{7,8})(?:\s*$|[,\s\]])')
PUBMED_URL = re.compile(r'pubmed\.ncbi\.nlm\.nih\.gov/(\d{7,8})', re.I)
YEAR_RE = re.compile(r'\b(19|20)\d{2}\b')
URL_RE = re.compile(r'https?://[^\s\)\]<>"]+')
# source-type label that Sonar sometimes puts in front of the citation
LABEL_PREFIX = re.compile(
    r'^\s*(guideline|systematic review|meta-?analysis|trial|rct|observational|'
    r'insufficient evidence)\s*[:\-–]\s*', re.I)


def load_answers():
    wb = openpyxl.load_workbook(os.path.join(RAW, KEY), data_only=True)
    ws = wb['Row Key']; h = [c.value for c in ws[1]]
    key = {d['Rating ID']: d for d in (dict(zip(h, r)) for r in ws.iter_rows(min_row=2, values_only=True)) if d.get('Rating ID')}
    wb2 = openpyxl.load_workbook(os.path.join(RAW, CONSENSUS), data_only=True)
    ws2 = wb2['Rating']; h2 = [c.value for c in ws2[1]]
    out = []
    for r in ws2.iter_rows(min_row=2, values_only=True):
        d = dict(zip(h2, r))
        if not d.get('Rating ID'):
            continue
        k = key[d['Rating ID']]
        out.append(dict(rating_id=d['Rating ID'], model=MODEL_NAME[k['Actual model']],
                        domain=d['Domain'], question_id=d['Question ID'], qtype=d['Question type'],
                        item_guideline=d['Guideline'], accuracy=d['Accuracy rating'],
                        glr=d['Guideline reference correct?'], text=d['Model answer'] or ''))
    return out


def is_bibliographic(inner):
    """A bracket expression counts as a citation if it contains a year and at least
    two commas. Bare references such as [1] or [Guideline] drop out."""
    s = LABEL_PREFIX.sub('', inner).strip()
    return bool(YEAR_RE.search(s)) and s.count(',') >= 2 and 20 <= len(s) <= 400


def parse_bracket(inner):
    s = re.sub(r'\s+', ' ', LABEL_PREFIX.sub('', inner)).strip()
    doi = DOI_RE.search(s)
    doi = doi.group(0).rstrip('.,;') if doi else ''
    rest = DOI_RE.sub('', s)
    m = PMID_LABEL.search(rest)
    if m:
        pmid = m.group(1)
    else:
        m2 = PMID_BARE.search(rest)
        pmid = m2.group(1) if m2 else ''
    ym = YEAR_RE.search(s)
    year = ym.group(0) if ym else ''
    author = s.split(',')[0].strip()
    tj = ''
    m3 = re.search(r'\b(?:19|20)\d{2}\b\s*,\s*(.+)', s)
    if m3:
        tj = DOI_RE.sub('', m3.group(1))
        tj = PMID_LABEL.sub('', tj)
        tj = re.sub(r',\s*\d{7,8}\s*$', '', tj)
        tj = re.sub(r'PMID[:\s]*not available[^,\]]*', '', tj, flags=re.I)
        tj = tj.strip().strip(',').strip()
    if '/' in tj:
        title, journal = tj.rsplit('/', 1)
    else:
        title, journal = tj, ''
    if author and YEAR_RE.fullmatch(author or ''):
        author = ''
    return dict(raw=s, author=author, year=year, title=title.strip(), journal=journal.strip(),
                pmid=pmid, doi=doi, url='', domain_host='', oe_tag='', source_format='bracket')


def parse_url_list(text):
    out, seen = [], set()
    for m in URL_RE.finditer(text):
        u = m.group(0).rstrip('.,;')
        if u in seen:
            continue
        seen.add(u)
        host = re.sub(r'^www\.', '', u.split('/')[2].lower())
        pm = PUBMED_URL.search(u)
        d = DOI_RE.search(u)
        out.append(dict(raw=u, author='', year='', title='', journal='',
                        pmid=pm.group(1) if pm else '', doi=d.group(0).rstrip('.,;') if d else '',
                        url=u, domain_host=host, oe_tag='', source_format='url_list'))
    return out


def doc_key(rec):
    """Identity of a document, in this order: DOI, PMID, normalised title."""
    if rec['doi']:
        return 'doi:' + rec['doi'].lower()
    if rec['pmid']:
        return 'pmid:' + rec['pmid']
    t = re.sub(r'[^a-z0-9 ]', ' ', (rec['title'] or '').lower())
    t = re.sub(r'\s+', ' ', t).strip()
    if len(t) >= 12:
        return 'title:' + t[:60]
    return 'url:' + rec['url'].lower() if rec['url'] else 'raw:' + rec['raw'][:60].lower()


def merge(records):
    """Merges duplicates within a response. Bibliographic data take precedence
    because the category rule has more information with them; the URL is kept
    in url_also."""
    by = collections.OrderedDict()
    for r in records:
        k = doc_key(r)
        if k not in by:
            by[k] = dict(r); by[k]['url_also'] = ''
            continue
        cur = by[k]
        if cur['source_format'] == 'url_list' and r['source_format'] == 'bracket':
            keep = dict(r); keep['url_also'] = cur['url']; keep['source_format'] = 'both'
            by[k] = keep
        elif cur['source_format'] == 'bracket' and r['source_format'] == 'url_list':
            cur['url_also'] = r['url']; cur['source_format'] = 'both'
    return list(by.values())


def main():
    answers = load_answers()
    import importlib.util
    spec = importlib.util.spec_from_file_location(
        'v1', os.path.join(ARCHIVE, 'extract_citations_v1.py'))
    v1 = importlib.util.module_from_spec(spec); spec.loader.exec_module(v1)

    BR = re.compile(r'\[([^\[\]]{20,400}?)\]')
    rows = []
    for a in answers:
        if a['model'] == 'Open Evidence':
            recs = [dict(r, source_format='numbered_list', url_also='') for r in v1.parse_oe(a["text"])]
        elif a['model'] == 'GPT-5.5':
            recs = [dict(r, source_format='bracket', url_also='') for r in v1.parse_gpt(a["text"])]
        else:
            urls = parse_url_list(a['text'])
            brs = [parse_bracket(m.group(1)) for m in BR.finditer(a['text']) if is_bibliographic(m.group(1))]
            recs = merge(urls + brs)
        for r in recs:
            r.setdefault('url_also', '')
            r.setdefault('source_format', '')
            r.setdefault('domain_host', r.get('domain', ''))
            rows.append(dict(rating_id=a['rating_id'], model=a['model'], domain=a['domain'],
                             question_id=a['question_id'], qtype=a['qtype'],
                             item_guideline=a['item_guideline'], accuracy=a['accuracy'], glr=a['glr'],
                             author=r.get('author', ''), year=r.get('year', ''), title=r.get('title', ''),
                             journal=r.get('journal', ''), pmid=r.get('pmid', ''), doi=r.get('doi', ''),
                             url=r.get('url', ''), domain_host=r.get('domain_host', ''),
                             oe_tag=r.get('oe_tag', ''), source_format=r['source_format'],
                             url_also=r['url_also'], raw=r.get('raw', '')))
    fn = os.path.join(OUT, 'citations_raw_v2.csv')
    with open(fn, 'w', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys())); w.writeheader(); w.writerows(rows)
    c = collections.Counter(r['model'] for r in rows)
    print('written:', fn)
    for m in ['Open Evidence', 'GPT-5.5', 'Sonar Reasoning Pro']:
        print('   %-22s %4d citations' % (m, c[m]))
    cf = collections.Counter(r['source_format'] for r in rows if r['model'] == 'Sonar Reasoning Pro')
    print('   Sonar by format:', dict(cf))
    print(REFERENCE)


REFERENCE = """
Reference values - a run must reproduce these:
  Open Evidence   431   GPT-5.5   168   Sonar Reasoning Pro   876
  Sonar by format: url_list 779, bracket 88, both 9
"""

if __name__ == '__main__':
    main()
